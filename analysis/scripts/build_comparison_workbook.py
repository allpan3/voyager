#!/usr/bin/env python3
"""Build one RTL-only CIM-vs-systolic workbook from a parsed sweep CSV.

Usage: build_comparison_workbook.py <summary_csv> <out_xlsx>

The workbook is intentionally per-layer. A hardware point is performance-
reportable only when every expected layer has a passing RTL row with the
required hardware counters. SystemC rows help identify attempted points but
never contribute cycles, utilization, comparisons, or stall measurements.
"""

import csv
import os
import sys
from collections import OrderedDict
from pathlib import Path

# Recover the declared Conda Python when vendor tools shadow python3 in PATH
try:
    import openpyxl
except ModuleNotFoundError:
    conda_python = Path(os.environ.get("CONDA_PREFIX", "")) / "bin" / "python"
    if conda_python.is_file() and conda_python.resolve() != Path(sys.executable).resolve():
        os.execv(str(conda_python), [str(conda_python), str(Path(__file__).resolve()), *sys.argv[1:]])
    raise

from openpyxl import Workbook
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

SRC, OUT = sys.argv[1], sys.argv[2]
rows = list(csv.DictReader(open(SRC)))
if not rows:
    sys.exit("summary CSV is empty")

FONT = "Arial"
NAVY = "17324D"
TEAL = "0F6B78"
TEAL_LIGHT = "D9EEF0"
SA_BLUE = "DCEAF7"
CIM_GREEN = "E2F0D9"
PASS_GREEN = "C6E0B4"
FAIL_RED = "F4CCCC"
GRAY = "667085"
WHITE = "FFFFFF"
GRID = "D0D5DD"

TITLE = Font(name=FONT, bold=True, size=15, color=WHITE)
HEADER = Font(name=FONT, bold=True, size=10, color=WHITE)
BODY = Font(name=FONT, size=10)
BODY_BOLD = Font(name=FONT, bold=True, size=10)
WRAP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(wrap_text=True, horizontal="center", vertical="center")
RIGHT = Alignment(horizontal="right", vertical="center")
THIN_BOTTOM = Border(bottom=Side(style="thin", color=GRID))

LAYER_SHORT = {
    "mobilebert_encoder_layer_0_ffn_0_output_dense_fused": "ffn_0_output_dense",
    "mobilebert_encoder_layer_0_output_bottleneck_dense_fused": "output_bottleneck_dense",
    "mobilebert_encoder_layer_0_attention_output_dense_fused": "attention_output_dense",
    "matmul_6_fused": "matmul_6 (attention context)",
    "matmul_2_fused": "matmul_2 (attention scores)",
    "layer1_0_conv1_fused": "ResNet-18 layer1.0 conv1",
    "layer2_0_conv1_fused": "ResNet-18 layer2.0 conv1",
    "layer4_1_conv2_fused": "ResNet-18 layer4.1 conv2",
}
REQUIRED_RTL_FIELDS = [
    "runtime_cycles",
    "matrix_unit_cycles",
    "ideal_cycles",
    "processor_active_cycles",
]
STALL_FIELDS = [
    ("input_unavailable_cycles", "Input unavailable"),
    ("input_backpressure_cycles", "Input backpressure"),
    ("weight_unavailable_cycles", "Weight unavailable"),
    ("weight_backpressure_cycles", "Weight backpressure"),
    ("result_backpressure_cycles", "Result backpressure"),
    ("accumulation_stall_cycles", "Accumulation stall"),
    ("output_fifo_full_cycles", "Output FIFO full"),
    ("output_backpressure_cycles", "Output backpressure"),
    ("mac_wait_weight_set_load_cycles", "MAC wait for weight-set load"),
    ("cim_completion_storage_stall_cycles", "CIM completion storage"),
    ("cim_result_slot_stall_cycles", "CIM result slots"),
    ("cim_completion_descriptor_stall_cycles", "CIM descriptor queue"),
]
STALL_DEFINITIONS = {
    "input_unavailable_cycles": (
        "Counted when the MatrixProcessor or CIMProcessor is ready to consume the next input/A beat, but "
        "InputController's window-buffer output has valid low.\n\n"
        "The beat may still be waiting for a memory response, input-buffer read, packing, transposition, or "
        "window formation. This counter identifies upstream input starvation but not which upstream stage caused it."
    ),
    "input_backpressure_cycles": (
        "SA: push_inputs has accepted an input beat and presents its decoded vector to the input skewer, but at least "
        "one per-row skewer FIFO cannot accept it. A PE row can stop consuming when its downstream input or partial-"
        "sum output cannot advance, when result-deskewer or accumulation-path pressure propagates backward through "
        "the partial-sum chain, or when a swap_weights input waits for the PE's one-entry next-weight FIFO.\n\n"
        "CIM: issue_operations has accepted an input beat and presents a fully formed MAC request, but CIMArray has "
        "ready low. This is the union of completion-storage unavailability and at least one selected tile having MAC "
        "ready low. Waiting for a B-set-ready token happens before the input is accepted and is counted separately as "
        "MAC wait for weight-set load."
    ),
    "weight_unavailable_cycles": (
        "Counted when the processor's weight loader is ready for the next weight beat, but the upstream weight "
        "channel has valid low.\n\n"
        "SA may be waiting for WeightController, weight-buffer fill/read, memory response, packing, or transposition. "
        "CIM streams directly from WeightController and may be waiting for memory response, packing, transposition, "
        "or production of the next resident-set row beat."
    ),
    "weight_backpressure_cycles": (
        "Counted when the upstream weight channel presents a valid beat but the processor's weight loader is not "
        "ready to accept it. This is the producer-valid/consumer-not-ready counterpart to Weight unavailable and is "
        "measured at the same processor input for SA and CIM.\n\n"
        "SA receives this channel from the weight-buffer output. A common cause is push_weights holding a previously "
        "accepted row while the serialized weight skewer or PE next-weight FIFOs drain. CIM receives the channel "
        "directly from WeightController. It can block before accepting the next beat while waiting for resident-bank "
        "reuse or while a previously accepted beat reaches CIMArray. MAC wait for weight-set load instead measures "
        "compute-side "
        "waiting for a resident set to finish loading."
    ),
    "result_backpressure_cycles": (
        "A raw array result is valid but the processor-side result consumer is not ready.\n\n"
        "Raw means the array's partial-sum output before bias addition, cross-tile accumulation, and final write-back. "
        "SA observes the psum output skewer. CIM observes CIMArray's CBeat result channel; it blocks when "
        "collect_results cannot forward the beat into the accumulation pipeline."
    ),
    "accumulation_stall_cycles": (
        "Union of accumulation-path blocking conditions for both SA and CIM: an accumulation-buffer read address is "
        "valid but not accepted; the accumulator expects read data that is not valid; a write request is valid but "
        "not accepted; or the accumulation-to-writeback FIFO cannot accept another completed sum.\n\n"
        "The component conditions can overlap, so this counter is not additive with other stalls."
    ),
    "output_fifo_full_cycles": (
        "Producer-side final-output blockage. write_back has a completed final accumulation and asserts enqueue "
        "valid, but the eight-entry final-output FIFO has no available slot.\n\n"
        "The FIFO drains through the processor output channel into MatrixUnit OutputController."
    ),
    "output_backpressure_cycles": (
        "Consumer-side final-output blockage. The final-output FIFO presents a valid entry on the processor output "
        "channel, but MatrixUnit OutputController is not ready to pop it.\n\n"
        "OutputController serializes the value and either sends it to memory or pushes it to the vector unit. This "
        "differs from Output FIFO full, which is measured at the FIFO's enqueue side."
    ),
    "mac_wait_weight_set_load_cycles": (
        "CIM only. The MAC issue loop is ready to begin the next operation group but the required B-set-ready token "
        "is unavailable because the loader has not finished that resident set. This is compute-side waiting for a B "
        "set, not weight-channel backpressure. The loader may still be receiving weight beats or writing them into "
        "CIMArray."
    ),
    "cim_completion_storage_stall_cycles": (
        "CIM only. A valid MAC request cannot reserve all completion resources. This is the union of descriptor-queue "
        "exhaustion and insufficient result slots in any selected output lane, including same-edge CBeat release "
        "credit. The component conditions can overlap, so their counters are diagnostic breakdowns and are not "
        "additive."
    ),
    "cim_result_slot_stall_cycles": (
        "CIM only. A valid MAC request lacks enough free result slots in at least one selected output lane. Reduced "
        "requests reserve one slot per selected lane; unreduced requests reserve one slot per input-axis tile in each "
        "selected lane. Slots remain reserved until the request's final CBeat is accepted."
    ),
    "cim_completion_descriptor_stall_cycles": (
        "CIM only. A valid MAC request finds every central completion-descriptor entry reserved and no descriptor is "
        "released by a final CBeat on the same cycle. Each accepted request owns one descriptor until its final CBeat "
        "is accepted."
    ),
}


# Convert a CSV value to an integer when present
def integer(value):
    if value in ("", None):
        return None
    return int(float(value))


# Interpret the parser's Boolean text
def passed(value):
    return str(value).lower() == "true"


# Return the stable identity of one hardware instance
def point_key(row):
    latency = row.get("cim_mac_latency", "") if row["backend"] == "cim" else ""
    return (
        row["config"],
        row["ic_port_width_bits"],
        row["oc_port_width_bits"],
        latency,
        row.get("clock_period_ns", ""),
    )


# Return a human-readable explicit port-width label
def port_label(row):
    ic = row["ic_port_width_bits"]
    oc = row["oc_port_width_bits"]
    return f"{ic}-bit" if ic == oc else f"{ic}/{oc}-bit"


# Return a complete instance label with explicit CIM macro geometry
def instance_label(row):
    base = f"{'CIM' if row['backend'] == 'cim' else 'SA'} {row['geometry']} | {port_label(row)} ports"
    if row["backend"] == "cim":
        mode = row.get("cim_mode") or "bit-parallel"
        sets = row.get("cim_b_sets") or "2"
        native_width = row.get("cim_macro_native_width") or "unknown"
        ch_in = row.get("cim_ch_in") or "?"
        ch_out = row.get("cim_ch_out") or "?"
        base += (
            f" | native {native_width} | ci{ch_in}/co{ch_out} | {mode} | {sets} sets"
            f" | MAC latency {row.get('cim_mac_latency') or '1'}"
            f" | {row.get('cim_result_slots_per_output_lane') or '?'} slots/lane"
        )
    return base


# Return the configured array MAC count
def macs(row):
    return integer(row["macs"]) or 0


# Apply the common title band
def style_title(sheet, title, last_column):
    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=last_column)
    cell = sheet.cell(1, 1, title)
    cell.font = TITLE
    cell.fill = PatternFill("solid", fgColor=NAVY)
    cell.alignment = Alignment(vertical="center")
    sheet.row_dimensions[1].height = 30
    sheet.sheet_view.showGridLines = False


# Apply the common table header
def style_header(sheet, row, first_column, last_column):
    for column in range(first_column, last_column + 1):
        cell = sheet.cell(row, column)
        cell.font = HEADER
        cell.fill = PatternFill("solid", fgColor=TEAL)
        cell.alignment = CENTER
    sheet.row_dimensions[row].height = 38


# Apply the common body-row styling
def style_body_row(sheet, row, first_column, last_column, fill):
    for column in range(first_column, last_column + 1):
        cell = sheet.cell(row, column)
        cell.font = BODY
        cell.fill = PatternFill("solid", fgColor=fill)
        cell.border = THIN_BOTTOM
        cell.alignment = WRAP if column == first_column else RIGHT


# Add a filterable Excel table with a stable style
def add_table(sheet, name, first_row, last_row, first_column, last_column):
    if last_row < first_row:
        return
    ref = (
        f"{get_column_letter(first_column)}{first_row}:"
        f"{get_column_letter(last_column)}{last_row}"
    )
    table = Table(displayName=name, ref=ref)
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    sheet.add_table(table)


# Assign explicit bounded column widths
def set_widths(sheet, widths):
    for index, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(index)].width = width


all_layers = list(OrderedDict((row["layer"], None) for row in rows))
workload_path = Path(SRC).resolve().with_name("workloads.csv")
if workload_path.is_file():
    with workload_path.open(newline="") as stream:
        declared_layers = [row["layer"] for row in csv.DictReader(stream)]
    layers = [layer for layer in declared_layers if layer in all_layers]
else:
    layers = all_layers


# Return whether one RTL row carries complete matrix performance evidence
def valid_rtl(row):
    return (
        row is not None
        and passed(row["passed"])
        and all(row.get(field) not in ("", None) for field in REQUIRED_RTL_FIELDS)
    )

point_rows = OrderedDict()
for row in rows:
    point_rows.setdefault(point_key(row), []).append(row)

points = []
for key, observations in point_rows.items():
    representative = observations[0]
    rtl = {row["layer"]: row for row in observations if row["sim"] == "rtl"}
    passing_layers = [
        layer
        for layer in layers
        if valid_rtl(rtl.get(layer))
    ]
    reportable = len(passing_layers) == len(layers)
    if reportable:
        reason = ""
    elif not rtl:
        reason = "RTL generation or simulation did not produce layer results"
    else:
        reason = f"{len(passing_layers)}/{len(layers)} workloads passed with complete RTL counters"
    points.append(
        {
            "key": key,
            "row": representative,
            "rtl": rtl,
            "reportable": reportable,
            "passing_layers": passing_layers,
            "reason": reason,
        }
    )

points.sort(
    key=lambda point: (
        macs(point["row"]),
        point["row"]["backend"] != "systolic",
        integer(point["row"]["ic_port_width_bits"]) or 0,
        integer(point["row"].get("cim_mac_latency")) or 0,
    )
)
reportable_points = [point for point in points if point["reportable"]]
sa_points = [point for point in points if point["row"]["backend"] == "systolic"]

if not sa_points:
    sys.exit("no systolic design point found")

baseline_sa = min(sa_points, key=lambda point: macs(point["row"]))


# Locate one point by backend, geometry, ports, and optional CIM latency
def find_point(backend, geometry, ports, latency=None):
    for point in points:
        row = point["row"]
        if row["backend"] != backend or row["geometry"] != geometry:
            continue
        if (row["ic_port_width_bits"], row["oc_port_width_bits"]) != ports:
            continue
        if backend == "cim" and str(row.get("cim_mac_latency") or "1") != str(latency or 1):
            continue
        return point
    return None


# Select the closest available SA reference without hiding unmatched CIM points
def reference_sa(point):
    row = point["row"]
    if row["backend"] == "systolic":
        return point
    ports = (row["ic_port_width_bits"], row["oc_port_width_bits"])
    return find_point("systolic", row["geometry"], ports) or baseline_sa

wb = Workbook()
wb.remove(wb.active)

# -------------------------------------------------------- Comparison Matrix
ws = wb.create_sheet("Comparisons")
config_headers = [
    "Config",
    "Backend",
    "Array",
    "Input port bits",
    "Output port bits",
    "Native width",
    "CIM mode",
    "Sets",
    "CH_IN",
    "CH_OUT",
    "MAC latency",
    "Result slots/lane",
    "SA reference",
]
metric_headers = []
for _ in layers:
    metric_headers.extend(["RTL cycles", "Speedup vs SA", "Utilization"])
comparison_headers = config_headers + metric_headers
comparison_last_column = len(comparison_headers)

style_title(
    ws,
    "Workload comparison matrix — one row per hardware point",
    comparison_last_column,
)
ws["A2"] = (
    "Filter the configuration columns to compare SA and CIM implementations directly. "
    "Each workload block shows RTL cycles, speedup against the closest available SA reference, and utilization; "
    "lower cycles and speedup above 1.00x are better."
)
ws["A2"].font = Font(name=FONT, italic=True, color=GRAY)
ws.merge_cells(
    start_row=2,
    start_column=1,
    end_row=2,
    end_column=comparison_last_column,
)

ws.merge_cells(
    start_row=4,
    start_column=1,
    end_row=4,
    end_column=len(config_headers),
)
ws.cell(4, 1, "Configuration")
metric_column = len(config_headers) + 1
for layer in layers:
    ws.merge_cells(
        start_row=4,
        start_column=metric_column,
        end_row=4,
        end_column=metric_column + 2,
    )
    ws.cell(4, metric_column, LAYER_SHORT.get(layer, layer))
    metric_column += 3
style_header(ws, 4, 1, comparison_last_column)

for column, heading in enumerate(comparison_headers, 1):
    ws.cell(5, column, heading)
style_header(ws, 5, 1, comparison_last_column)

comparison_start = 6
comparison_row = comparison_start
result_row_map = {}
pending_matrix_formulas = []
matrix_points = sorted(
    reportable_points,
    key=lambda point: (
        point["row"]["backend"] != "systolic",
        macs(point["row"]),
        integer(point["row"]["ic_port_width_bits"]) or 0,
        int((point["row"].get("cim_macro_native_width") or "0b").rstrip("b")),
        point["row"].get("cim_mode") or "",
        integer(point["row"].get("cim_b_sets")) or 0,
        integer(point["row"].get("cim_ch_in")) or 0,
        integer(point["row"].get("cim_ch_out")) or 0,
        integer(point["row"].get("cim_result_slots_per_output_lane")) or 0,
    ),
)
for point in matrix_points:
    row = point["row"]
    reference = reference_sa(point)
    values = [
        row["config"],
        "CIM" if row["backend"] == "cim" else "SA",
        row["geometry"],
        integer(row["ic_port_width_bits"]),
        integer(row["oc_port_width_bits"]),
        int((row.get("cim_macro_native_width") or "0b").rstrip("b"))
        if row["backend"] == "cim"
        else None,
        row.get("cim_mode") if row["backend"] == "cim" else None,
        integer(row.get("cim_b_sets")) if row["backend"] == "cim" else None,
        integer(row.get("cim_ch_in")) if row["backend"] == "cim" else None,
        integer(row.get("cim_ch_out")) if row["backend"] == "cim" else None,
        integer(row.get("cim_mac_latency"))
        if row["backend"] == "cim"
        else None,
        integer(row.get("cim_result_slots_per_output_lane"))
        if row["backend"] == "cim"
        else None,
        instance_label(reference["row"]),
    ]
    for column, value in enumerate(values, 1):
        ws.cell(comparison_row, column, value)
    fill = CIM_GREEN if row["backend"] == "cim" else SA_BLUE
    style_body_row(ws, comparison_row, 1, comparison_last_column, fill)
    for layer_index, layer in enumerate(layers):
        first_metric_column = len(config_headers) + 1 + layer_index * 3
        pending_matrix_formulas.append(
            (
                comparison_row,
                point["key"],
                reference["key"],
                layer,
                first_metric_column,
            )
        )
    comparison_row += 1

comparison_last = comparison_row - 1
ws.auto_filter.ref = (
    f"A5:{get_column_letter(comparison_last_column)}{comparison_last}"
)
ws.freeze_panes = f"{get_column_letter(len(config_headers) + 1)}6"
set_widths(
    ws,
    [31, 9, 10, 14, 15, 13, 13, 8, 9, 9, 12, 18, 34]
    + [14, 17, 14] * len(layers),
)

# --------------------------------------------------------- Per-Layer Results
ws = wb.create_sheet("Per-Layer Results")
style_title(ws, "Passing RTL results — one row per instance and layer", 15)
ws["A2"] = (
    "MatrixUnit utilization = ideal logical MAC cycles / MatrixUnit start-to-done cycles. "
    "Failed hardware points are excluded and appear only in Design Points."
)
ws["A2"].font = Font(name=FONT, italic=True, color=GRAY)
ws.merge_cells("A2:O2")
result_headers = [
    "Instance",
    "Config",
    "Backend",
    "Array",
    "Input port bits",
    "Output port bits",
    "CIM MAC latency",
    "Layer",
    "RTL cycles",
    "MatrixUnit cycles",
    "Raw array-slot ideal cycles",
    "Logical A width",
    "Logical MAC issue cycles",
    "Ideal logical MAC cycles",
    "MatrixUnit utilization",
]
for column, heading in enumerate(result_headers, 1):
    ws.cell(4, column, heading)
style_header(ws, 4, 1, len(result_headers))

result_row = 5
for point in reportable_points:
    row = point["row"]
    fill = SA_BLUE if row["backend"] == "systolic" else CIM_GREEN
    for layer in all_layers:
        rtl = point["rtl"].get(layer)
        if not valid_rtl(rtl):
            continue
        values = [
            instance_label(row),
            row["config"],
            row["backend"],
            row["geometry"],
            integer(row["ic_port_width_bits"]),
            integer(row["oc_port_width_bits"]),
            integer(row.get("cim_mac_latency")) if row["backend"] == "cim" else None,
            LAYER_SHORT.get(layer, layer),
            integer(rtl["runtime_cycles"]),
            integer(rtl["matrix_unit_cycles"]),
            integer(rtl["raw_ideal_cycles"]),
            integer(rtl["logical_a_width_bits"]),
            integer(rtl["logical_mac_issue_cycles"]),
            integer(rtl["ideal_cycles"]),
        ]
        for column, value in enumerate(values, 1):
            ws.cell(result_row, column, value)
        ws.cell(result_row, 15, f"=IFERROR(N{result_row}/J{result_row},\"\")")
        ws.cell(result_row, 15).number_format = "0.0%"
        for column in range(5, 15):
            ws.cell(result_row, column).number_format = "#,##0"
        ws.cell(result_row, 13).number_format = "0x"
        style_body_row(ws, result_row, 1, len(result_headers), fill)
        result_row_map[(point["key"], layer)] = result_row
        result_row += 1
add_table(
    ws,
    "PerLayerRtlResults",
    4,
    result_row - 1,
    1,
    len(result_headers),
)
ws.freeze_panes = "I5"
set_widths(ws, [43, 31, 11, 11, 14, 15, 16, 29, 12, 19, 20, 15, 23, 23, 18])

# Link matrix formulas to the auditable per-layer results
comparison_ws = wb["Comparisons"]
for row, point_key_value, reference_key, layer, first_column in pending_matrix_formulas:
    point_result_row = result_row_map[(point_key_value, layer)]
    reference_result_row = result_row_map[(reference_key, layer)]
    cycle_column = get_column_letter(first_column)
    comparison_ws.cell(
        row,
        first_column,
        f"='Per-Layer Results'!I{point_result_row}",
    )
    comparison_ws.cell(
        row,
        first_column + 1,
        f'=IFERROR(\'Per-Layer Results\'!I{reference_result_row}/{cycle_column}{row},"")',
    )
    comparison_ws.cell(
        row,
        first_column + 2,
        f"='Per-Layer Results'!O{point_result_row}",
    )
    comparison_ws.cell(row, first_column).number_format = "#,##0"
    comparison_ws.cell(row, first_column + 1).number_format = "0.00x"
    comparison_ws.cell(row, first_column + 2).number_format = "0.0%"

for layer_index, _ in enumerate(layers):
    first_column = len(config_headers) + 1 + layer_index * 3
    cycle_letter = get_column_letter(first_column)
    speedup_letter = get_column_letter(first_column + 1)
    cycle_range = (
        f"{cycle_letter}{comparison_start}:{cycle_letter}{comparison_last}"
    )
    speedup_range = (
        f"{speedup_letter}{comparison_start}:{speedup_letter}{comparison_last}"
    )
    comparison_ws.conditional_formatting.add(
        cycle_range,
        ColorScaleRule(
            start_type="min",
            start_color=PASS_GREEN,
            end_type="max",
            end_color=FAIL_RED,
        ),
    )
    comparison_ws.conditional_formatting.add(
        speedup_range,
        ColorScaleRule(
            start_type="min",
            start_color=FAIL_RED,
            mid_type="num",
            mid_value=1,
            mid_color="FFF2CC",
            end_type="max",
            end_color=PASS_GREEN,
        ),
    )

# ------------------------------------------------------------- Stall Analysis
ws = wb.create_sheet("Stall Analysis")
stall_headers = [
    "Instance",
    "Layer",
    "Processor active cycles",
    *[heading for _, heading in STALL_FIELDS],
]
style_title(ws, "Passing RTL stall counters — SA and CIM", len(stall_headers))
ws["A2"] = (
    "Counters are simultaneous conditions and may overlap. "
    "Complete signal-level definitions are in the final Methodology sheet."
)
ws["A2"].font = Font(name=FONT, italic=True, color=GRAY)
ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(stall_headers))
for column, heading in enumerate(stall_headers, 1):
    ws.cell(4, column, heading)
style_header(ws, 4, 1, len(stall_headers))

stall_row = 5
for point in reportable_points:
    row = point["row"]
    fill = SA_BLUE if row["backend"] == "systolic" else CIM_GREEN
    for layer in all_layers:
        rtl = point["rtl"].get(layer)
        if not valid_rtl(rtl):
            continue
        values = [
            instance_label(row),
            LAYER_SHORT.get(layer, layer),
            integer(rtl["processor_active_cycles"]),
            *[
                integer(rtl.get(field))
                if row["backend"] == "cim" or not field.startswith("cim_")
                else None
                for field, _ in STALL_FIELDS
            ],
        ]
        for column, value in enumerate(values, 1):
            ws.cell(stall_row, column, value)
            if column >= 3:
                ws.cell(stall_row, column).number_format = "#,##0"
        style_body_row(ws, stall_row, 1, len(stall_headers), fill)
        stall_row += 1
add_table(
    ws,
    "RtlStallAnalysis",
    4,
    stall_row - 1,
    1,
    len(stall_headers),
)
ws.freeze_panes = "D5"
set_widths(ws, [43, 29, 19] + [18] * len(STALL_FIELDS))

# -------------------------------------------------------------- Design Points
ws = wb.create_sheet("Design Points")
design_headers = [
    "Instance",
    "Config",
    "Backend",
    "Technology",
    "Datatype",
    "Array",
    "Input port bits",
    "Output port bits",
    "Clock period (ns)",
    "CIM MAC latency",
    "CIM mode",
    "CIM B sets",
    "Macro native width",
    "CH_IN",
    "CH_OUT",
    "Result slots/lane",
    "RTL status",
    "Passing RTL workloads",
    "Failure reason",
]
style_title(ws, "Sweep design points and RTL status", len(design_headers))
ws["A2"] = (
    "MAC latency is a CIM-only parameter. SA cells are intentionally blank. "
    "A failed point has no performance or comparison rows."
)
ws["A2"].font = Font(name=FONT, italic=True, color=GRAY)
ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(design_headers))
for column, heading in enumerate(design_headers, 1):
    ws.cell(4, column, heading)
style_header(ws, 4, 1, len(design_headers))
design_row = 5
for point in points:
    row = point["row"]
    values = [
        instance_label(row),
        row["config"],
        row["backend"],
        row.get("technology") or None,
        row.get("datatype") or None,
        row["geometry"],
        integer(row["ic_port_width_bits"]),
        integer(row["oc_port_width_bits"]),
        float(row["clock_period_ns"]),
        integer(row.get("cim_mac_latency")) if row["backend"] == "cim" else None,
        row.get("cim_mode") if row["backend"] == "cim" else None,
        integer(row.get("cim_b_sets")) if row["backend"] == "cim" else None,
        row.get("cim_macro_native_width") or None,
        integer(row.get("cim_ch_in")),
        integer(row.get("cim_ch_out")),
        integer(row.get("cim_result_slots_per_output_lane"))
        if row["backend"] == "cim" else None,
        "PASS" if point["reportable"] else "FAILED",
        f"{len(point['passing_layers'])}/{len(layers)}",
        point["reason"],
    ]
    for column, value in enumerate(values, 1):
        ws.cell(design_row, column, value)
    fill = SA_BLUE if row["backend"] == "systolic" else CIM_GREEN
    style_body_row(ws, design_row, 1, len(design_headers), fill)
    status_column = design_headers.index("RTL status") + 1
    ws.cell(design_row, status_column).fill = PatternFill(
        "solid", fgColor=PASS_GREEN if point["reportable"] else FAIL_RED
    )
    ws.cell(design_row, status_column).font = BODY_BOLD
    design_row += 1
add_table(
    ws,
    "SweepDesignPoints",
    4,
    design_row - 1,
    1,
    len(design_headers),
)
ws.freeze_panes = "F5"
set_widths(
    ws,
    [43, 31, 11, 13, 12, 11, 14, 15, 16, 16, 16, 12, 18, 10, 10,
     18, 12, 19, 54],
)

# --------------------------------------------------------------- Methodology
ws = wb.create_sheet("Methodology")
style_title(ws, "Methodology", 3)
method_rows = [
    (
        "MatrixUnit cycles",
        "Harness Matrix Unit Runtime / clock period",
        "Measured from the MatrixUnit start handshake through the MatrixUnit done handshake. It includes "
        "OutputController drain and excludes any later vector-unit tail. The clock period comes from the sweep "
        "manifest.",
    ),
    (
        "Raw array-slot ideal cycles",
        "Logged MAC work / physical array slots",
        "The simulator's raw total-MACs / array-slots value. It assumes one operation per physical slot per cycle. "
        "Multiply it by the configured logical-MAC issue interval for narrow or bit-serial CIM macros.",
    ),
    (
        "Logical A width",
        "Matched input port bits / array K",
        "The activation operand width is derived from the configured matched port and array geometry, independent of "
        "the datatype name.",
    ),
    (
        "Logical MAC issue cycles",
        "Derived from logical A width and CIM base A/B/C widths, CH_IN, and mode",
        "Bit-parallel uses ceil(logical A width / base A width). Bit-serial uses the exact padded slice-walking "
        "interval implemented by CIMElement. SA uses one cycle. No datatype names are enumerated.",
    ),
    (
        "Ideal logical MAC cycles",
        "Raw array-slot ideal cycles * logical MAC issue cycles",
        "Minimum full-array cycles for the layer's logical MAC work after accounting for the configured CIM issue "
        "window.",
    ),
    (
        "MatrixUnit utilization",
        "Ideal logical MAC cycles / MatrixUnit cycles",
        "The workbook's logical-MAC utilization metric for matrix operations.",
    ),
    (
        "Processor active cycles",
        "Hardware processor_active_cycles counter",
        "Internal diagnostic window from processor parameter acceptance through processor write-back completion. "
        "Retained on Stall Analysis but not used for utilization.",
    ),
    (
        "Stall counters",
        "SA and CIM",
        "Stall conditions may overlap and must not be added. Complete signal-level definitions follow below.",
    ),
    (
        "CIM result slots per output lane",
        "CIM_ARRAY_RESULT_SLOTS_PER_OUTPUT_LANE",
        "Physical tile-vector result entries owned independently by each output lane. Production reduced requests "
        "reserve one slot in every selected lane until their final C beat is accepted; standalone unreduced requests "
        "reserve INPUT_AXIS_TILES slots per selected lane.",
    ),
    (
        "Result-slot selection",
        "Explicit sweep parameter",
        "The workbook does not infer that any slot count is sufficient. Use the measured CIM result-slot and "
        "completion-storage stall counters to evaluate each hardware and workload combination.",
    ),
]
method_rows.extend(
    [
        *[
            (
                heading,
                "CIM only" if field.startswith("cim_") else "SA and CIM",
                STALL_DEFINITIONS[field].replace("\n\n", " "),
            )
            for field, heading in STALL_FIELDS
        ],
    ]
)
for column, heading in enumerate(["Field", "Definition", "Notes"], 1):
    ws.cell(3, column, heading)
style_header(ws, 3, 1, 3)
for row_index, values in enumerate(method_rows, 4):
    for column, value in enumerate(values, 1):
        ws.cell(row_index, column, value)
        ws.cell(row_index, column).font = BODY_BOLD if column == 1 else BODY
        ws.cell(row_index, column).alignment = WRAP
        ws.cell(row_index, column).border = THIN_BOTTOM
    ws.cell(row_index, 1).fill = PatternFill("solid", fgColor=TEAL_LIGHT)
    ws.row_dimensions[row_index].height = max(34, 15 * (1 + len(str(values[2])) // 82))
add_table(ws, "DefinitionTable", 3, 3 + len(method_rows), 1, 3)
set_widths(ws, [28, 38, 92])

for sheet in wb.worksheets:
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_view.zoomScale = 90

wb.calculation.fullCalcOnLoad = True
wb.calculation.forceFullCalc = True
wb.calculation.calcMode = "auto"
wb.save(OUT)

print(
    f"saved {OUT}: {len(wb.sheetnames)} sheets, "
    f"{len(points)} hardware points, {len(reportable_points)} RTL-pass points, "
    f"{len(matrix_points)} comparison rows, {len(layers)} workloads"
)
